#!/usr/bin/python3
import os
from typing import Dict, List
from docx import Document
import random
import tkinter as tk
from tkinter import messagebox


def process_data(file_name):
    data = {}
    with open(file_name, "r") as file:
        key = None
        value = ""
        for line in file:
            line = line.strip()
            if line.endswith(":"):
                if key is not None:
                    data[key] = value.strip()
                    value = ""
                key = line[:-1]
            else:
                value += line + "\n"
        if key is not None:
            data[key] = value.strip()
    return data


def is_valid_pair(lunch, dinner):
    valid_pairs = {"A": ["A", "B", "C", "D"], "B": ["A"], "C": ["A"], "D": ["A"]}
    return dinner in valid_pairs.get(lunch, [])


def generate_plan(breakfast_data: Dict, lunch_data: Dict, dinner_data: Dict):
    diet_plan = []
    for week in range(4):
        week_plan = []
        used_options = set()
        for day in range(7):
            lunch = random.choice(list(lunch_data.keys()))
            dinner_options = [
                opt
                for opt in dinner_data.keys()
                if is_valid_pair(lunch, opt) and opt != lunch
            ]
            if not dinner_options:
                continue
            dinner = random.choice(dinner_options)

            if day in [5, 6]:  # Saturday and Sunday
                breakfast_options = breakfast_data["A"]
            elif day in [1, 2, 4]:  # Tuesday, Wednesday, Friday
                breakfast_options = breakfast_data["B"]
            else:  # Monday, Thursday
                breakfast_options = breakfast_data["C"]

            if day in [0, 1, 2, 3, 4]:
                spuntino_mattutino = breakfast_data["SPUNTINO"]
            else:
                spuntino_mattutino = ""

            while (breakfast_options, lunch, dinner) in used_options:
                lunch = random.choice(list(lunch_data.keys()))
                dinner_options = [
                    opt
                    for opt in dinner_data.keys()
                    if is_valid_pair(lunch, opt) and opt != lunch
                ]
                if not dinner_options:
                    continue
                dinner = random.choice(dinner_options)

            week_plan.append(
                (
                    breakfast_options,
                    spuntino_mattutino,
                    lunch_data[lunch],
                    lunch_data["SPUNTINO"],
                    dinner_data[dinner],
                )
            )

            used_options.add((breakfast_options, lunch, dinner))

        diet_plan.append(week_plan)
    return diet_plan


def save_plan_txt(diet_plan: List):
    desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
    file_path = os.path.join(desktop_path, "diet_plan.txt")

    with open(file_path, "w", encoding="ansi") as output_file:
        output_file.write("Piano Dietetico(4 settimane):\n\n")
        for i, week in enumerate(diet_plan, 1):
            output_file.write(f"Settimana {i}:\n")
            for day, meal in enumerate(week, 1):
                output_file.write(f"Giorno {day}:\n")
                output_file.write(f"Colazione: \n{meal[0]}\n\n")
                output_file.write(f"Spuntino Mattutino: \n{meal[1]}\n\n")
                output_file.write(f"Pranzo: \n{meal[2]}\n\n")
                output_file.write(f"Spuntino Pomeridiano: \n{meal[3]}\n\n")
                output_file.write(f"Cena: \n{meal[4]}\n\n")
            output_file.write("\n")

    abs_file_path = os.path.abspath(file_path)
    messagebox.showinfo(
        "Salvataggio .txt completato",
        f"Il piano dietetico e' stato salvato nel file {abs_file_path}",
    )


def save_plan_docx(diet_plan: List):
    doc = Document()
    doc.add_heading("Piano Dietetico (4 settimane)", level=1)

    for i, week in enumerate(diet_plan, 1):
        doc.add_heading(f"Settimana {i}", level=2)
        for day, meal in enumerate(week, 1):
            doc.add_heading(f"Giorno {day}", level=3)
            doc.add_heading("Colazione:", level=4)
            doc.add_paragraph(meal[0])
            doc.add_heading("Spuntino Mattutino:", level=4)
            doc.add_paragraph(meal[1])
            doc.add_heading("Pranzo:", level=4)
            doc.add_paragraph(meal[2])
            doc.add_heading("Spuntino Pomeridiano:", level=4)
            doc.add_paragraph(meal[3])
            doc.add_heading("Cena:", level=4)
            doc.add_paragraph(meal[4])

    desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
    file_path = os.path.join(desktop_path, "diet_plan.docx")
    doc.save(file_path)
    abs_file_path = os.path.abspath(file_path)
    messagebox.showinfo(
        "Salvataggio .docx completato",
        f"Il piano dietetico e' stato salvato nel file {abs_file_path}",
    )


def show_gui(diet_plan: List):
    root = tk.Tk()
    root.title("Generatore Piano Dietetico")
    root.resizable(False, False)

    title = tk.Label(
        root,
        text="Generatore Piano Dietetico",
        font=("Arial", 16, "bold"),
    )
    title.pack(padx=20, pady=20)

    button_txt = tk.Button(
        root,
        text="Salva il piano dietetico da 4 settimane come file di testo(.txt)",
        font=("Arial", 16),
        command=lambda: save_plan_txt(diet_plan),
    )
    button_txt.pack(padx=20, pady=20)

    button_docx = tk.Button(
        root,
        text="Salva il piano dietetico da 4 settimane come file word(.docx)",
        font=("Arial", 16),
        command=lambda: save_plan_docx(diet_plan),
    )
    button_docx.pack(padx=10, pady=10)

    root.mainloop()


def main():
    breakfast = "breakfast.txt"
    lunch = "lunch.txt"
    dinner = "dinner.txt"

    breakfast_data = process_data(breakfast)
    lunch_data = process_data(lunch)
    dinner_data = process_data(dinner)

    diet_plan = generate_plan(breakfast_data, lunch_data, dinner_data)
    print([elem for row in diet_plan for elem in row])
    show_gui(diet_plan)


if __name__ == "__main__":
    main()
